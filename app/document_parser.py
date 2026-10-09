from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class MaterialError(ValueError):
    pass


def extract_text(filename: str, raw: bytes, max_chars: int) -> str:
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".pdf":
            text = _extract_pdf(raw)
        elif extension == ".docx":
            text = _extract_docx(raw)
        elif extension == ".txt":
            text = _extract_txt(raw)
        else:
            raise MaterialError("Formatos permitidos: PDF, DOCX y TXT.")
    except MaterialError:
        raise
    except Exception as exc:
        raise MaterialError("No se pudo leer el documento. Revisa que no esté dañado o protegido.") from exc

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not text:
        raise MaterialError(
            "No se encontró texto legible. Los PDF escaneados necesitan OCR antes de subirlos."
        )
    if len(text) > max_chars:
        raise MaterialError(
            f"El documento supera el máximo de {max_chars:,} caracteres extraídos. Sube un capítulo más corto."
        )
    return text


def _extract_pdf(raw: bytes) -> str:
    reader = PdfReader(BytesIO(raw), strict=False)
    pages = []
    for page_number, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(f"[Página {page_number}]\n{page_text}")
    return "\n\n".join(pages)


def _extract_docx(raw: bytes) -> str:
    document = Document(BytesIO(raw))
    blocks = [
        f"[Sección {index}] {paragraph.text}"
        for index, paragraph in enumerate(document.paragraphs, start=1)
        if paragraph.text.strip()
    ]
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(f"[Tabla {table_index}]\n" + "\n".join(rows))
    return "\n\n".join(blocks)


def _extract_txt(raw: bytes) -> str:
    content = raw.decode("utf-8-sig")
    lines = [line.strip() for line in content.splitlines() if line.strip()]
    return "\n".join(f"[Línea {index}] {line}" for index, line in enumerate(lines, start=1))
